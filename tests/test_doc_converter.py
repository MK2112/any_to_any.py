import os
import docx
import shutil
import pytest
import tempfile
import numpy as np
import pymupdf as pypdf

from utils.category import Category
from core.converter.doc_converter import DocumentConverter
from tests.test_fixtures import controller_instance, setup_file_handler_mock
from unittest.mock import Mock, patch, mock_open, MagicMock


@pytest.fixture
def mock_file_handler():
    handler = Mock()
    handler.join_back.return_value = "/mock/path/file.ext"
    handler.post_process.return_value = None
    handler.has_visuals.return_value = True
    setup_file_handler_mock(handler)
    return handler


@pytest.fixture
def mock_prog_logger():
    return Mock()


@pytest.fixture
def mock_event_logger():
    logger = Mock()
    logger.info.return_value = None
    return logger


@pytest.fixture
def document_converter(mock_file_handler, mock_prog_logger, mock_event_logger):
    # DocumentConverter instance creation with mocked dependencies
    return DocumentConverter(
        file_handler=mock_file_handler,
        prog_logger=mock_prog_logger,
        event_logger=mock_event_logger,
        locale="English",
    )


@pytest.fixture
def temp_output_dir():
    # Temp directory for test outputs
    temp_dir = tempfile.mkdtemp()
    yield temp_dir
    shutil.rmtree(temp_dir, ignore_errors=True)


@pytest.fixture
def sample_file_paths():
    return {
        Category.DOCUMENT: [
            ("/path", "test_doc", "docx"),
            ("/path", "test_pdf", "pdf"),
            ("/path", "test_srt", "srt"),
            ("/path", "test_pptx", "pptx"),
        ],
        Category.IMAGE: [
            ("/path", "test_image", "jpg"),
            ("/path", "test_gif", "gif"),
            ("/path", "test_png", "png"),
        ],
        Category.MOVIE: [
            ("/path", "test_movie", "mp4"),
            ("/path", "test_video", "avi"),
        ],
    }


class TestDocumentConverterInit:
    def test_init_with_default_locale(
        self, mock_file_handler, mock_prog_logger, mock_event_logger
    ):
        converter = DocumentConverter(
            file_handler=mock_file_handler,
            prog_logger=mock_prog_logger,
            event_logger=mock_event_logger,
        )
        assert converter.locale == "English"

    def test_init_with_custom_locale(
        self, mock_file_handler, mock_prog_logger, mock_event_logger
    ):
        converter = DocumentConverter(
            file_handler=mock_file_handler,
            prog_logger=mock_prog_logger,
            event_logger=mock_event_logger,
            locale="Spanish",
        )
        assert converter.locale == "Spanish"


class TestDocumentConverterToMarkdown:
    @patch("core.converter.doc_converter.markdownify")
    def test_to_markdown_docx_conversion(
        self, mock_markdownify, document_converter, temp_output_dir, sample_file_paths
    ):
        # Test DOCX to Markdown conversion
        docx_path = os.path.join(temp_output_dir, "test_doc.docx")
        doc = docx.Document()
        doc.add_paragraph("Test content")
        doc.save(docx_path)

        mock_markdownify.return_value = "# Test content"
        document_converter.file_handler.join_back.return_value = docx_path
        document_converter.to_markdown(
            output=temp_output_dir,
            file_paths={Category.DOCUMENT: [("path", "test_doc", "docx")]},
            format="md",
            delete=False,
        )

        mock_markdownify.assert_called_once()

        output_path = os.path.join(temp_output_dir, "test_doc.md")

        assert os.path.exists(output_path)
        with open(output_path, "r") as f:
            content = f.read()
            assert content == "# Test content"

        document_converter.file_handler.post_process.assert_called_once()

    @patch("core.converter.doc_converter.platform.system", return_value="Windows")
    def test_docx_pdf_uses_reportlab_on_windows(
        self,
        _mock_platform,
        document_converter,
        temp_output_dir,
    ):
        docx_path = os.path.join(temp_output_dir, "test_doc.docx")
        doc = docx.Document()
        doc.add_paragraph("Test content")
        doc.save(docx_path)

        document_converter.file_handler.join_back.return_value = docx_path
        with patch.object(
            document_converter, "_docx_to_pdf_reportlab"
        ) as reportlab_mock:
            document_converter.to_pdf(
                output=temp_output_dir,
                file_paths={
                    Category.DOCUMENT: [("/path", "test_doc", "docx")],
                    Category.IMAGE: [],
                    Category.MOVIE: [],
                },
                format="pdf",
                delete=False,
            )

        reportlab_mock.assert_called_once()

    @patch("core.converter.doc_converter.platform.system", return_value="Linux")
    def test_docx_pdf_falls_back_when_weasyprint_missing(
        self,
        _mock_platform,
        document_converter,
        temp_output_dir,
    ):
        docx_path = os.path.join(temp_output_dir, "test_doc.docx")
        doc = docx.Document()
        doc.add_paragraph("Test content")
        doc.save(docx_path)

        document_converter.file_handler.join_back.return_value = docx_path
        with patch("core.converter.doc_converter.HTML", None):
            with patch.object(
                document_converter, "_docx_to_pdf_reportlab"
            ) as reportlab_mock:
                document_converter.to_pdf(
                    output=temp_output_dir,
                    file_paths={
                        Category.DOCUMENT: [("/path", "test_doc", "docx")],
                        Category.IMAGE: [],
                        Category.MOVIE: [],
                    },
                    format="pdf",
                    delete=False,
                )

        reportlab_mock.assert_called_once()

    @patch("mammoth.convert_to_html")
    @patch("os.makedirs")
    def test_to_markdown_with_images(
        self, mock_makedirs, mock_mammoth, document_converter, temp_output_dir
    ):
        # Setup mock for mammoth with image handling
        mock_image = Mock()
        mock_image.content_type = "image/png"
        mock_image.open.return_value.__enter__ = Mock()
        mock_image.open.return_value.__exit__ = Mock(return_value=None)
        mock_image.open.return_value.__enter__.return_value.read.return_value = (
            b"image_data"
        )

        mock_document = Mock()
        mock_document.value = "<p>Content with image</p>"
        mock_mammoth.return_value = mock_document

        docx_files = {Category.DOCUMENT: [("/path", "test_doc", "docx")]}

        with (
            patch("builtins.open", mock_open()),
            patch("markdownify.markdownify", return_value="markdown"),
        ):
            document_converter.to_markdown(temp_output_dir, docx_files, "md", False)

        mock_makedirs.assert_called()
        mock_mammoth.assert_called_once()

    def test_to_markdown_image_links_are_relative(self, tmp_path):
        import re
        import logging

        from PIL import Image
        from core.utils.file_handler import FileHandler

        Image.new("RGB", (8, 8), "red").save(tmp_path / "pic.png")
        document = docx.Document()
        document.add_picture(str(tmp_path / "pic.png"))
        document.save(tmp_path / "report.docx")

        out = tmp_path / "out"
        out.mkdir()
        file_handler = FileHandler(logging.getLogger("md-relative"), "English")
        converter = DocumentConverter(file_handler, MagicMock(), MagicMock())
        converter.to_markdown(
            str(out),
            {
                Category.IMAGE: [],
                Category.MOVIE: [],
                Category.DOCUMENT: [(str(tmp_path) + os.sep, "report", "docx")],
            },
            "md",
            False,
        )

        md_path = out / "report.md"
        assert md_path.exists()
        markdown = md_path.read_text(encoding="utf-8")
        match = re.search(r"!\[[^\]]*\]\(([^)]+)\)", markdown)
        assert match, markdown
        src = match.group(1)
        assert not os.path.isabs(src)
        assert (md_path.parent / src).is_file()


class TestDocumentConverterToPdf:
    @patch("core.converter.doc_converter.pypdf.open")
    @patch("core.converter.doc_converter.pypdf.Pixmap")
    def test_to_pdf_image_conversion(
        self,
        mock_pixmap,
        mock_pypdf_open,
        document_converter,
        temp_output_dir,
        sample_file_paths,
    ):
        mock_doc = MagicMock()
        mock_page = MagicMock()
        mock_pix = MagicMock(width=100, height=100)
        mock_pixmap.return_value = mock_pix
        mock_doc.new_page.return_value = mock_page
        mock_pypdf_open.return_value = mock_doc

        document_converter.file_handler.join_back.return_value = os.path.join(
            "path", "test_image.jpg"
        )

        document_converter.to_pdf(
            output=temp_output_dir,
            file_paths={
                Category.IMAGE: [("path", "test_image", "jpg")],
                Category.MOVIE: [],
                Category.DOCUMENT: [],
            },
            format="pdf",
            delete=False,
        )

        mock_doc.save.assert_called_once()
        mock_doc.close.assert_called()

    @patch("core.converter.doc_converter.gif_to_frames")
    @patch("core.converter.doc_converter.pypdf.open")
    @patch("os.listdir")
    @patch("os.path.join")
    @patch("core.converter.doc_converter.pypdf.Pixmap")
    @patch("shutil.rmtree")
    def test_to_pdf_gif_conversion(
        self,
        mock_rmtree,
        mock_pixmap,
        mock_join,
        mock_listdir,
        mock_pypdf_open,
        mock_gif_to_frames,
        document_converter,
        temp_output_dir,
    ):
        mock_doc = MagicMock()
        mock_page = MagicMock()
        mock_pix = MagicMock(width=100, height=100)
        mock_pixmap.return_value = mock_pix
        mock_doc.new_page.return_value = mock_page
        mock_pypdf_open.return_value = mock_doc
        mock_listdir.return_value = ["frame1.png", "frame2.png"]
        mock_join.side_effect = lambda *args: "/".join(args)

        def mock_gif_to_frames_impl(output, file_paths, file_handler):
            gif_dir = os.path.join(output, file_paths[Category.IMAGE][0][1])
            os.makedirs(gif_dir, exist_ok=True)
            with open(os.path.join(gif_dir, "frame1.png"), "wb") as f:
                f.write(b"fake image data")
            with open(os.path.join(gif_dir, "frame2.png"), "wb") as f:
                f.write(b"fake image data")

        mock_gif_to_frames.side_effect = mock_gif_to_frames_impl

        document_converter.file_handler.join_back.return_value = os.path.join(
            "path", "test_gif.gif"
        )

        document_converter.to_pdf(
            output=temp_output_dir,
            file_paths={
                Category.IMAGE: [["path", "test_gif", "gif"]],
                Category.MOVIE: [],
                Category.DOCUMENT: [],
            },
            format="pdf",
            delete=False,
        )

        # Verify the PDF was created and cleanup was performed
        mock_doc.save.assert_called_once()
        mock_rmtree.assert_called()

    @patch("PIL.Image.fromarray")
    @patch("os.remove")
    @patch("core.converter.doc_converter.VideoFileClip")
    @patch("core.converter.doc_converter.pypdf.open")
    @patch("core.converter.doc_converter.pypdf.Pixmap")
    @patch("os.path.join")
    def test_to_pdf_movie_conversion(
        self,
        mock_join,
        mock_pixmap,
        mock_pypdf_open,
        mock_video_clip,
        mock_remove,
        mock_image_fromarray,
        document_converter,
        temp_output_dir,
    ):
        mock_doc = MagicMock()
        mock_page = MagicMock()
        mock_pix = MagicMock(width=100, height=100)
        mock_pixmap.return_value = mock_pix
        mock_doc.new_page.return_value = mock_page
        mock_pypdf_open.return_value = mock_doc

        mock_video = MagicMock()
        mock_video.fps = 30
        mock_video.duration = 2.0

        frame = np.zeros((100, 100, 3), dtype=np.uint8)
        mock_video.iter_frames.return_value = [frame, frame]  # Two frames
        mock_video_clip.return_value.__enter__.return_value = mock_video
        mock_img = MagicMock()
        mock_image_fromarray.return_value = mock_img

        document_converter.file_handler.join_back.return_value = os.path.join(
            "path", "test_movie.mp4"
        )
        mock_join.side_effect = lambda *args: "/".join(args)

        document_converter.to_pdf(
            output=temp_output_dir,
            file_paths={
                Category.MOVIE: [["path", "test_movie", "mp4"]],
                Category.IMAGE: [],
                Category.DOCUMENT: [],
            },
            format="pdf",
            delete=False,
        )

        # Verify the PDF was created
        mock_doc.save.assert_called_once()
        mock_video_clip.assert_called_once()


class TestDocumentConverterToSubtitles:
    @patch("subprocess.run")
    @patch("os.path.exists")
    @patch("os.path.getsize")
    @patch("utils.language_support.get_translation")
    def test_to_subtitles_success(
        self,
        mock_get_translation,
        mock_getsize,
        mock_exists,
        mock_subprocess,
        document_converter,
        temp_output_dir,
    ):
        mock_exists.return_value = True
        mock_getsize.return_value = 100  # Non-empty file
        mock_get_translation.return_value = "Extracting subtitles"
        mock_subprocess.return_value = Mock()

        movie_files = {
            Category.MOVIE: [("/path", "test_movie", "mp4")],
            Category.IMAGE: [],
            Category.DOCUMENT: [],
        }

        document_converter.to_subtitles(temp_output_dir, movie_files, "srt", False)
        mock_subprocess.assert_called()
        document_converter.file_handler.post_process.assert_called()

    @patch("subprocess.run")
    @patch("os.path.exists")
    @patch("os.path.getsize")
    @patch("utils.language_support.get_translation")
    def test_to_subtitles_fallback_method(
        self,
        mock_get_translation,
        mock_getsize,
        mock_exists,
        mock_subprocess,
        document_converter,
        temp_output_dir,
    ):
        # Setup mocks for first attempt failure, second attempt success
        mock_exists.side_effect = [False, True]  # First fails, second succeeds
        mock_getsize.side_effect = [0, 100]  # First empty, second has content
        mock_get_translation.return_value = "Extracting subtitles"
        mock_subprocess.return_value = Mock()

        movie_files = {
            Category.MOVIE: [("/path", "test_movie", "mp4")],
            Category.IMAGE: [],
            Category.DOCUMENT: [],
        }

        document_converter.to_subtitles(temp_output_dir, movie_files, "srt", False)
        assert mock_subprocess.call_count == 2  # Both attempts called


class TestDocumentConverterEdgeCases:
    def test_empty_file_paths(self, document_converter, temp_output_dir):
        empty_paths = {Category.DOCUMENT: [], Category.IMAGE: [], Category.MOVIE: []}
        document_converter.to_markdown(temp_output_dir, empty_paths, "md", False)
        document_converter.to_pdf(temp_output_dir, empty_paths, "pdf", False)
        document_converter.to_subtitles(temp_output_dir, empty_paths, "srt", False)
        document_converter.to_office(temp_output_dir, empty_paths, "docx", False)

    @patch("pymupdf.open")
    def test_pdf_skip_existing_pdf(
        self, mock_pypdf_open, document_converter, temp_output_dir
    ):
        pdf_files = {
            Category.DOCUMENT: [("/path", "test", "pdf")],
            Category.IMAGE: [],
            Category.MOVIE: [],
        }
        with patch("core.converter.image_converter.gif_to_frames"):
            document_converter.to_pdf(temp_output_dir, pdf_files, "pdf", False)
        mock_pypdf_open.assert_not_called()


class TestDocumentConverterPdfSplitting:
    def test_split_pdf_with_real_file(self, document_converter, temp_output_dir):
        # Test PDF splitting with a real 5 page PDF file
        pdf_path = os.path.join(temp_output_dir, "test.pdf")
        doc = pypdf.open()

        for i in range(5):
            page = doc.new_page()
            page.insert_text((72, 72), f"Page {i + 1}")
        doc.save(pdf_path)
        doc.close()

        # Test splitting into two parts: pages 1-2 and 3-5
        doc_path_set = (temp_output_dir, "test", "pdf")
        document_converter.file_handler.join_back.return_value = pdf_path

        # Split the PDF
        document_converter.split_pdf(
            output=temp_output_dir,
            doc_path_set=doc_path_set,
            page_ranges="1-2,3-5",
            delete=False,
            format="pdf",
        )

        # Verify the output files were created with the correct naming pattern
        part1_path = os.path.join(temp_output_dir, "test_split_1_1-2.pdf")
        part2_path = os.path.join(temp_output_dir, "test_split_2_3-5.pdf")

        assert os.path.exists(part1_path)
        assert os.path.exists(part2_path)

        # Verify the page counts are correct
        part1 = pypdf.open(part1_path)
        part2 = pypdf.open(part2_path)

        assert len(part1) == 2  # First part has 2 pages
        assert len(part2) == 3  # Second part has 3 pages

        # Clean up
        part1.close()
        part2.close()

        # Verify the content of the first page of each part
        part1 = pypdf.open(part1_path)
        part2 = pypdf.open(part2_path)

        assert "Page 1" in part1[0].get_text()
        assert "Page 3" in part2[0].get_text()

        part1.close()
        part2.close()

    def test_split_pdf_all_splits_every_page(
        self, document_converter, temp_output_dir
    ):
        pdf_path = os.path.join(temp_output_dir, "test.pdf")
        doc = pypdf.open()
        for i in range(3):
            doc.new_page().insert_text((72, 72), f"Page {i + 1}")
        doc.save(pdf_path)
        doc.close()

        document_converter.file_handler.join_back.return_value = pdf_path
        document_converter.split_pdf(
            output=temp_output_dir,
            doc_path_set=(temp_output_dir, "test", "pdf"),
            page_ranges="all",
            delete=False,
            format="pdf",
        )

        for page in range(1, 4):
            part_path = os.path.join(
                temp_output_dir, f"test_split_{page}_{page}-{page}.pdf"
            )
            assert os.path.exists(part_path)
            part = pypdf.open(part_path)
            assert len(part) == 1
            assert f"Page {page}" in part[0].get_text()
            part.close()

    def test_split_pdf_rest_keeps_whole_document(
        self, document_converter, temp_output_dir
    ):
        pdf_path = os.path.join(temp_output_dir, "test.pdf")
        doc = pypdf.open()
        for i in range(3):
            doc.new_page().insert_text((72, 72), f"Page {i + 1}")
        doc.save(pdf_path)
        doc.close()

        document_converter.file_handler.join_back.return_value = pdf_path
        document_converter.split_pdf(
            output=temp_output_dir,
            doc_path_set=(temp_output_dir, "test", "pdf"),
            page_ranges="rest",
            delete=False,
            format="pdf",
        )

        part_path = os.path.join(temp_output_dir, "test_1-3.pdf")
        assert os.path.exists(part_path)
        part = pypdf.open(part_path)
        assert len(part) == 3
        part.close()


class TestDocumentConverterSplitWarnings:
    def _make_pdf(self, path, pages):
        doc = pypdf.open()
        for i in range(pages):
            doc.new_page().insert_text((72, 72), f"Page {i + 1}")
        doc.save(path)
        doc.close()

    def test_split_pdf_invalid_ranges_warns_and_produces_nothing(
        self, document_converter, mock_event_logger, temp_output_dir
    ):
        pdf_path = os.path.join(temp_output_dir, "test.pdf")
        self._make_pdf(pdf_path, 3)
        document_converter.file_handler.join_back.return_value = pdf_path

        document_converter.split_pdf(
            output=temp_output_dir,
            doc_path_set=(temp_output_dir, "test", "pdf"),
            page_ranges="999",
            delete=False,
            format="pdf",
        )

        mock_event_logger.warning.assert_called()
        assert os.listdir(temp_output_dir) == ["test.pdf"]

    def test_split_pdf_multi_range_post_processes_all_outputs(
        self, document_converter, temp_output_dir
    ):
        pdf_path = os.path.join(temp_output_dir, "test.pdf")
        self._make_pdf(pdf_path, 5)
        document_converter.file_handler.join_back.return_value = pdf_path

        document_converter.split_pdf(
            output=temp_output_dir,
            doc_path_set=(temp_output_dir, "test", "pdf"),
            page_ranges="1-2,3-5",
            delete=False,
            format="pdf",
        )

        assert document_converter.file_handler.post_process.call_count == 2


def test_controller_split_non_pdf_warns(controller_instance, caplog):
    with caplog.at_level("WARNING"):
        controller_instance.split(
            {Category.DOCUMENT: [("/tmp/", "notes", "docx")]},
            page_ranges="1-2",
        )
    assert "skipping" in caplog.text.lower()


def test_run_with_pdf_format_and_split_actually_splits(
    controller_instance, tmp_path
):
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    pdf_path = input_dir / "doc.pdf"
    doc = pypdf.open()
    for i in range(4):
        doc.new_page().insert_text((72, 72), f"Page {i + 1}")
    doc.save(str(pdf_path))
    doc.close()

    output_dir = tmp_path / "output"
    output_dir.mkdir()

    controller_instance.run(
        [str(input_dir)],
        format=["pdf"],
        output=str(output_dir),
        framerate=None,
        quality=None,
        split="1-2",
        merge=False,
        concat=False,
        delete=False,
        across=False,
        recursive=False,
        dropzone=False,
        language=None,
        workers=1,
    )

    out_path = output_dir / "doc_1-2.pdf"
    assert out_path.exists()
    part = pypdf.open(str(out_path))
    assert len(part) == 2
    part.close()


class TestDocumentConverterIntegration:
    # Integration tests that test multiple components together
    @patch("core.converter.image_converter.gif_to_frames")
    @patch("pymupdf.open")
    @patch("mammoth.convert_to_html")
    @patch("weasyprint.HTML")
    def test_docx_to_pdf_integration(
        self,
        mock_weasyprint,
        mock_mammoth,
        mock_pypdf_open,
        mock_gif_to_frames,
        document_converter,
        temp_output_dir,
    ):
        # Test DOCX to PDF conversion integration
        mock_document = Mock()
        mock_document.value = "<p>Test content</p>"
        mock_mammoth.return_value = mock_document

        mock_html = Mock()
        mock_weasyprint.return_value = mock_html

        docx_files = {
            Category.DOCUMENT: [("/path", "test_doc", "docx")],
            Category.IMAGE: [],
            Category.MOVIE: [],
        }

        with patch("builtins.open", mock_open()):
            document_converter.to_pdf(temp_output_dir, docx_files, "pdf", False)

        mock_mammoth.assert_called_once()
        document_converter.file_handler.post_process.assert_called()


@pytest.mark.parametrize(
    "page_ranges,total,expected",
    [
        (None, 10, None),
        ("", 5, None),
        ("all", 7, [(1, 1), (2, 2), (3, 3), (4, 4), (5, 5), (6, 6), (7, 7)]),
        ("rest", 7, [(1, 7)]),
        ("3", 10, [(3, 3)]),
        ("1-3", 10, [(1, 3)]),
        ("2-5,7-8", 10, [(2, 5), (7, 8)]),
        ("2-5;7-8", 10, [(2, 5), (7, 8)]),
        ("12-end", 20, [(12, 20)]),
        ("2-end,4-6", 10, [(2, 10), (4, 6)]),
        ("3-4,rest", 6, [(3, 4), (5, 6)]),
        ("2-4,2-4", 8, [(2, 4)]),
        ("999", 5, None),
        ("a-b", 5, None),
    ],
)
def test_parse_page_ranges_various(document_converter, page_ranges, total, expected):
    res = document_converter._parse_page_ranges(page_ranges, total)
    assert res == expected


def test_parse_page_ranges_rest_after_range(document_converter):
    # 'rest' after a specific range should pick up after highest end
    res = document_converter._parse_page_ranges("1-2,rest", 5)
    assert res == [(1, 2), (3, 5)]


class TestDocumentConverterGifFrameOrder:
    def test_to_pdf_gif_frames_processed_in_numeric_order(
        self, document_converter, temp_output_dir
    ):
        frame_dir = os.path.join(temp_output_dir, "anim")
        os.makedirs(frame_dir)
        for name in ("anim-10.png", "anim-2.png", "anim-1.png"):
            open(os.path.join(frame_dir, name), "wb").close()

        with patch(
            "core.converter.doc_converter.gif_to_frames"
        ) as mock_gif_to_frames, patch(
            "core.converter.doc_converter.pypdf.open"
        ) as mock_pypdf_open, patch(
            "core.converter.doc_converter.pypdf.Pixmap"
        ) as mock_pixmap:
            mock_gif_to_frames.return_value = {"anim": frame_dir}
            mock_pypdf_open.return_value = MagicMock()
            mock_pixmap.return_value = MagicMock(width=100, height=100)
            document_converter.to_pdf(
                temp_output_dir,
                {
                    Category.IMAGE: [("dir", "anim", "gif")],
                    Category.MOVIE: [],
                    Category.DOCUMENT: [],
                },
                "pdf",
                False,
            )

        seen = [c.args[0] for c in mock_pixmap.call_args_list]
        assert seen == [
            os.path.join(frame_dir, "anim-1.png"),
            os.path.join(frame_dir, "anim-2.png"),
            os.path.join(frame_dir, "anim-10.png"),
        ]


class TestDocumentConverterPptxTitle:
    def test_pptx_titleless_slide_does_not_crash(self, tmp_path):
        from pptx import Presentation

        from core.utils.file_handler import FileHandler

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        assert slide.shapes.title is None
        prs.save(str(tmp_path / "deck.pptx"))
        out = tmp_path / "out"
        out.mkdir()
        file_handler = FileHandler(event_logger=MagicMock(), locale="English")
        converter = DocumentConverter(file_handler, MagicMock(), MagicMock())
        converter.to_office(
            str(out),
            {
                Category.IMAGE: [],
                Category.MOVIE: [],
                Category.DOCUMENT: [(str(tmp_path) + "/", "deck", "pptx")],
            },
            "docx",
            False,
        )

        assert (out / "deck.docx").exists()


class TestDocumentConverterSrtEncoding:
    def test_srt_to_pdf_non_ascii_under_ascii_locale(self, tmp_path):
        # True regression test: under an ASCII locale, bare open() raises
        # UnicodeDecodeError on non-ASCII subtitles while explicit utf-8-sig
        # reads them fine.
        import subprocess
        import sys
        import textwrap

        repo_root = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
        (tmp_path / "subs.srt").write_bytes(
            "1\n00:00:01,000 --> 00:00:02,000\nHéllo wörld\n\n".encode("utf-8")
        )
        out = tmp_path / "out"
        out.mkdir()
        doc_tuple = repr([(str(tmp_path) + "/", "subs", "srt")])
        script = textwrap.dedent(
            f"""
            import sys
            sys.path.insert(0, {repo_root!r})
            from unittest.mock import MagicMock
            from utils.category import Category
            from core.converter.doc_converter import DocumentConverter
            from core.utils.file_handler import FileHandler
            fh = FileHandler(event_logger=MagicMock(), locale="English")
            conv = DocumentConverter(fh, MagicMock(), MagicMock())
            conv.to_pdf(
                {str(out)!r},
                {{Category.IMAGE: [], Category.MOVIE: [],
                  Category.DOCUMENT: {doc_tuple}}},
                "pdf",
                False,
            )
            print("OK")
            """
        )
        env = dict(os.environ)
        env.update(
            {
                "PYTHONUTF8": "0",
                "LC_ALL": "C",
                "LANG": "C",
                "LANGUAGE": "C",
                "PYTHONCOERCECLOCALE": "0",
            }
        )
        proc = subprocess.run(
            [sys.executable, "-c", script],
            capture_output=True,
            text=True,
            env=env,
            timeout=120,
        )
        assert proc.returncode == 0, proc.stderr
        assert (out / "subs.pdf").exists()


def test_parse_page_ranges_mixed_delimiters(document_converter):
    assert document_converter._parse_page_ranges("1-2,3-4;5-6", 10) == [
        (1, 2),
        (3, 4),
        (5, 6),
    ]


class TestOfficePlainTextSources:

    @staticmethod
    def _converter():
        from core.utils.file_handler import FileHandler

        file_handler = FileHandler(event_logger=MagicMock(), locale="English")
        return DocumentConverter(file_handler, MagicMock(), MagicMock())

    @staticmethod
    def _paths(tmp_path, name, ext):
        return {
            Category.IMAGE: [],
            Category.MOVIE: [],
            Category.DOCUMENT: [(str(tmp_path) + os.sep, name, ext)],
        }

    def test_srt_to_docx_carries_text_over(self, tmp_path):
        (tmp_path / "subs.srt").write_text(
            "1\n00:00:01,000 --> 00:00:02,000\nHello world\n",
            encoding="utf-8",
        )
        out = tmp_path / "out"
        out.mkdir()

        self._converter().to_office(
            str(out), self._paths(tmp_path, "subs", "srt"), "docx", True
        )

        produced = out / "subs.docx"
        assert produced.exists()
        paragraphs = [p.text for p in docx.Document(str(produced)).paragraphs]
        assert "Hello world" in paragraphs
        # A real DOCX was produced, so -d may remove the source
        assert not (tmp_path / "subs.srt").exists()

    def test_md_to_docx_carries_text_over(self, tmp_path):
        (tmp_path / "notes.md").write_text("# Title\nSome text\n", encoding="utf-8")
        out = tmp_path / "out"
        out.mkdir()

        self._converter().to_office(
            str(out), self._paths(tmp_path, "notes", "md"), "docx", True
        )

        produced = out / "notes.docx"
        assert produced.exists()
        paragraphs = [p.text for p in docx.Document(str(produced)).paragraphs]
        assert "# Title" in paragraphs
        assert "Some text" in paragraphs
        assert not (tmp_path / "notes.md").exists()

    @pytest.mark.parametrize("name,ext", [("empty", "srt"), ("empty", "md")])
    def test_empty_text_source_is_skipped_and_kept(self, tmp_path, name, ext):
        (tmp_path / f"{name}.{ext}").write_text("", encoding="utf-8")
        out = tmp_path / "out"
        out.mkdir()

        self._converter().to_office(
            str(out), self._paths(tmp_path, name, ext), "docx", True
        )

        assert list(out.glob("*.docx")) == []
        assert (tmp_path / f"{name}.{ext}").exists()
